"""
Script de exportación a Excel (.xlsx) para el equipo investigador.
Genera un archivo con múltiples hojas: resumen, datos crudos, frecuencias.

Uso:
    python scripts/exportar_excel.py
    python scripts/exportar_excel.py --output resultados_mayo.xlsx
"""
import sys
import os
import argparse
from datetime import datetime, timezone

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app import create_app, db
from app.models import Participante, HabitosDigitales, ConsecuenciasVividas
from sqlalchemy import func

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("❌ Instala openpyxl: pip install openpyxl")
    sys.exit(1)


# ── Estilos ────────────────────────────────────────────────────────────────────
AZUL = "1A3A5C"
ACENTO = "E8A020"
BLANCO = "FFFFFF"
GRIS = "F4F6F9"

def estilo_header(cell, bg=AZUL, fg=BLANCO, bold=True):
    cell.font = Font(bold=bold, color=fg, name="Calibri", size=10)
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def estilo_subheader(cell):
    estilo_header(cell, bg="2B5F8E", fg=BLANCO)

def auto_width(ws, min_w=10, max_w=50):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_w), max_w)


# ── Hoja 1: Datos crudos ───────────────────────────────────────────────────────
def hoja_datos_crudos(wb, participantes):
    ws = wb.create_sheet("Datos crudos")

    headers = [
        "ID", "Fecha respuesta",
        # Demográficos
        "Edad", "Género", "Ciudad", "Departamento", "Nivel educativo", "Ocupación",
        # Hábitos
        "Horas/día internet", "Redes sociales",
        "Usa banca digital", "Usa e-commerce", "Usa apps salud", "Usa transporte", "Usa gov",
        "Comparte nombre", "Comparte foto", "Comparte ubicación",
        "Comparte teléfono", "Comparte financiero", "Comparte laboral",
        "Lee términos", "Conoce Ley 1581", "Preocupación privacidad (1-5)", "Cambió config privacidad",
        # Consecuencias
        "Exp. suplantación", "Exp. ciberacoso", "Exp. fraude", "Exp. pérdida privacidad",
        "Exp. robo cuenta", "Exp. spam", "Exp. discriminación laboral", "Sin experiencias negativas",
        "Impacto emocional", "Cambió comportamiento", "Descripción cambio",
        "Afectó relaciones", "Afectó empleo", "Afectó crédito",
        "Control datos (1-5)", "Confianza plataformas (1-5)",
    ]

    ws.append(headers)
    for cell in ws[1]:
        estilo_header(cell)
    ws.row_dimensions[1].height = 30

    for p in participantes:
        h = p.habitos
        c = p.consecuencias
        row = [
            p.id,
            p.fecha_respuesta.strftime("%Y-%m-%d %H:%M") if p.fecha_respuesta else "",
            p.edad, p.genero, p.ciudad, p.departamento, p.nivel_educativo, p.ocupacion,
            h.horas_diarias_internet if h else "", h.redes_sociales if h else "",
            "Sí" if h and h.usa_banca_digital else "No",
            "Sí" if h and h.usa_comercio_electronico else "No",
            "Sí" if h and h.usa_apps_salud else "No",
            "Sí" if h and h.usa_apps_transporte else "No",
            "Sí" if h and h.usa_servicios_gobierno else "No",
            "Sí" if h and h.comparte_nombre_real else "No",
            "Sí" if h and h.comparte_foto_perfil else "No",
            "Sí" if h and h.comparte_ubicacion else "No",
            "Sí" if h and h.comparte_telefono else "No",
            "Sí" if h and h.comparte_datos_financieros else "No",
            "Sí" if h and h.comparte_informacion_laboral else "No",
            h.lee_terminos_condiciones if h else "",
            "Sí" if h and h.conoce_ley_datos_colombia else "No",
            h.nivel_preocupacion_privacidad if h else "",
            "Sí" if h and h.ha_cambiado_config_privacidad else "No",
            "Sí" if c and c.experimento_suplantacion else "No",
            "Sí" if c and c.experimento_ciberacoso else "No",
            "Sí" if c and c.experimento_fraude_financiero else "No",
            "Sí" if c and c.experimento_perdida_privacidad else "No",
            "Sí" if c and c.experimento_robo_cuenta else "No",
            "Sí" if c and c.experimento_spam_publicidad else "No",
            "Sí" if c and c.experimento_discriminacion_laboral else "No",
            "Sí" if c and c.ninguna_experiencia_negativa else "No",
            c.impacto_emocional if c else "",
            "Sí" if c and c.cambio_comportamiento_digital else "No",
            c.descripcion_cambio if c and c.descripcion_cambio else "",
            "Sí" if c and c.afecto_relaciones_sociales else "No",
            "Sí" if c and c.afecto_oportunidades_laborales else "No",
            "Sí" if c and c.afecto_acceso_credito else "No",
            c.siente_control_datos if c else "",
            c.confianza_plataformas if c else "",
        ]
        ws.append(row)
        # Fila de datos con fondo alternado
        if ws.max_row % 2 == 0:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=GRIS)

    auto_width(ws)


# ── Hoja 2: Frecuencias ────────────────────────────────────────────────────────
def hoja_frecuencias(wb, participantes):
    ws = wb.create_sheet("Frecuencias")
    total = len(participantes)

    def pct(n):
        return round((n / total * 100), 1) if total > 0 else 0

    def seccion(titulo, fila_actual):
        ws.cell(row=fila_actual, column=1, value=titulo)
        cell = ws.cell(row=fila_actual, column=1)
        estilo_subheader(cell)
        ws.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=4)
        fila_actual += 1

        ws.cell(row=fila_actual, column=1, value="Categoría")
        ws.cell(row=fila_actual, column=2, value="N")
        ws.cell(row=fila_actual, column=3, value="%")
        for col in range(1, 4):
            estilo_header(ws.cell(row=fila_actual, column=col), bg="2B5F8E")
        return fila_actual + 1

    row = 1

    # Género
    row = seccion("GÉNERO", row)
    generos = {}
    for p in participantes:
        generos[p.genero] = generos.get(p.genero, 0) + 1
    for g, n in sorted(generos.items(), key=lambda x: -x[1]):
        ws.cell(row=row, column=1, value=g)
        ws.cell(row=row, column=2, value=n)
        ws.cell(row=row, column=3, value=pct(n))
        row += 1
    row += 1

    # Horas diarias
    row = seccion("HORAS DIARIAS DE INTERNET", row)
    orden_horas = ["menos_1h", "1_3h", "3_6h", "mas_6h"]
    horas_count = {}
    for p in participantes:
        if p.habitos:
            h = p.habitos.horas_diarias_internet
            horas_count[h] = horas_count.get(h, 0) + 1
    for clave in orden_horas:
        n = horas_count.get(clave, 0)
        ws.cell(row=row, column=1, value=clave.replace("_", " "))
        ws.cell(row=row, column=2, value=n)
        ws.cell(row=row, column=3, value=pct(n))
        row += 1
    row += 1

    # Experiencias negativas
    row = seccion("EXPERIENCIAS NEGATIVAS REPORTADAS", row)
    campos_exp = [
        ("Suplantación de identidad", "experimento_suplantacion"),
        ("Ciberacoso", "experimento_ciberacoso"),
        ("Fraude financiero", "experimento_fraude_financiero"),
        ("Pérdida de privacidad", "experimento_perdida_privacidad"),
        ("Robo de cuenta", "experimento_robo_cuenta"),
        ("Spam / publicidad no deseada", "experimento_spam_publicidad"),
        ("Discriminación laboral", "experimento_discriminacion_laboral"),
        ("Ninguna experiencia negativa", "ninguna_experiencia_negativa"),
    ]
    for label, campo in campos_exp:
        n = sum(1 for p in participantes if p.consecuencias and getattr(p.consecuencias, campo))
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=n)
        ws.cell(row=row, column=3, value=pct(n))
        row += 1
    row += 1

    # Impacto emocional
    row = seccion("IMPACTO EMOCIONAL", row)
    impactos = {}
    for p in participantes:
        if p.consecuencias:
            i = p.consecuencias.impacto_emocional
            impactos[i] = impactos.get(i, 0) + 1
    for imp, n in sorted(impactos.items(), key=lambda x: -x[1]):
        ws.cell(row=row, column=1, value=imp)
        ws.cell(row=row, column=2, value=n)
        ws.cell(row=row, column=3, value=pct(n))
        row += 1

    auto_width(ws)


# ── Hoja 3: Resumen ejecutivo ──────────────────────────────────────────────────
def hoja_resumen(wb, participantes):
    ws = wb.create_sheet("Resumen", 0)
    total = len(participantes)

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 20

    ws["A1"] = "INVESTIGACIÓN: Datos Personales e Internet · Colombia 2026"
    ws["A1"].font = Font(bold=True, size=13, color=AZUL, name="Calibri")
    ws.merge_cells("A1:B1")

    ws["A2"] = f"Generado el {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M')} UTC"
    ws["A2"].font = Font(italic=True, size=9, color="8898AA")
    ws.merge_cells("A2:B2")

    datos = [
        ("", ""),
        ("Total encuestas completadas", total),
        ("Edad promedio", round(sum(p.edad for p in participantes) / total, 1) if total else 0),
        ("", ""),
        ("Conocen Ley 1581 de 2012",
         f"{round(sum(1 for p in participantes if p.habitos and p.habitos.conoce_ley_datos_colombia) / total * 100, 1) if total else 0}%"),
        ("Cambiaron config. de privacidad",
         f"{round(sum(1 for p in participantes if p.habitos and p.habitos.ha_cambiado_config_privacidad) / total * 100, 1) if total else 0}%"),
        ("Preocupación promedio por privacidad (1-5)",
         round(sum(p.habitos.nivel_preocupacion_privacidad for p in participantes if p.habitos) / total, 2) if total else 0),
        ("", ""),
        ("Han vivido al menos una exp. negativa",
         f"{round(sum(1 for p in participantes if p.consecuencias and not p.consecuencias.ninguna_experiencia_negativa) / total * 100, 1) if total else 0}%"),
        ("Impacto emocional severo",
         f"{round(sum(1 for p in participantes if p.consecuencias and p.consecuencias.impacto_emocional == 'severo') / total * 100, 1) if total else 0}%"),
        ("Cambiaron comportamiento digital",
         f"{round(sum(1 for p in participantes if p.consecuencias and p.consecuencias.cambio_comportamiento_digital) / total * 100, 1) if total else 0}%"),
        ("Control promedio sobre sus datos (1-5)",
         round(sum(p.consecuencias.siente_control_datos for p in participantes if p.consecuencias) / total, 2) if total else 0),
        ("Confianza promedio en plataformas (1-5)",
         round(sum(p.consecuencias.confianza_plataformas for p in participantes if p.consecuencias) / total, 2) if total else 0),
    ]

    for i, (label, valor) in enumerate(datos, start=4):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=valor)
        if label:
            ws.cell(row=i, column=1).font = Font(name="Calibri", size=10)
            ws.cell(row=i, column=2).font = Font(bold=True, name="Calibri", size=10, color=AZUL)
        if i % 2 == 0 and label:
            ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=GRIS)
            ws.cell(row=i, column=2).fill = PatternFill("solid", fgColor=GRIS)


# ── Main ───────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="Exportar datos de investigación a Excel")
    parser.add_argument("--output", default=None, help="Nombre del archivo de salida")
    args = parser.parse_args()

    app = create_app()
    with app.app_context():
        participantes = (
            Participante.query
            .filter_by(completado=True)
            .join(HabitosDigitales)
            .join(ConsecuenciasVividas)
            .all()
        )

        if not participantes:
            print("⚠  No hay encuestas completadas para exportar.")
            return

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Quita la hoja vacía por defecto

        hoja_resumen(wb, participantes)
        hoja_datos_crudos(wb, participantes)
        hoja_frecuencias(wb, participantes)

        nombre = args.output or f"investigacion_digital_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        wb.save(nombre)
        print(f"✅ Archivo exportado: {nombre}")
        print(f"   · {len(participantes)} encuestas completas")
        print(f"   · Hojas: Resumen, Datos crudos, Frecuencias")


if __name__ == "__main__":
    main()